from __future__ import annotations

import argparse
import os
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


BLUE = colors.HexColor("#1F4E79")
BLUE_LIGHT = colors.HexColor("#D9EAF7")
BLUE_PALE = colors.HexColor("#EFF6FC")
GRAY = colors.HexColor("#666666")
GRAY_LIGHT = colors.HexColor("#F3F5F7")
YELLOW_PALE = colors.HexColor("#FFF6D9")
ORANGE = colors.HexColor("#C98200")
GREEN = colors.HexColor("#237A57")
RED = colors.HexColor("#B44B4B")

BASE_DIR = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = BASE_DIR / "input" / "category_difference_report(2).xlsx"
DEFAULT_OUTPUT = BASE_DIR / "output" / "pdf" / "full_category_performance_report_zh.pdf"


GROUPS = {
    "收入类": ["Wages", "Centrelink", "All Other Credits"],
    "支出类": [
        "Information",
        "Donations",
        "Education",
        "Home Improvement",
        "Insurance",
        "Subscription TV",
        "Pet Care",
        "Entertainment",
        "Retail",
        "Utilities",
        "Gyms and other memberships",
        "Personal Care",
        "Rent",
        "Health",
        "Travel",
        "Groceries",
        "Automotive",
        "Gambling",
        "Department Stores",
        "Telecommunications",
        "Dining Out",
        "Fees",
        "Transport",
    ],
    "负债类": [
        "SACC Loans",
        "Non SACC Loans",
        "Credit Card Repayments",
        "Debt Consolidation",
        "Debt Collection",
        "Dishonours",
        "Overdrawn",
        "Unknown Loans",
    ],
    "转账专项": ["External Transfers", "Internal Transfer"],
}

GROUP_OF = {category: group for group, categories in GROUPS.items() for category in categories}
FOCUS_CATEGORIES = {
    "External Transfers",
    "Internal Transfer",
    "Groceries",
    "Dining Out",
    "Retail",
    "Gambling",
    "All Other Credits",
    "Wages",
    "Rent",
    "SACC Loans",
    "Non SACC Loans",
}


def register_fonts() -> None:
    regular = Path(r"C:\Windows\Fonts\Deng.ttf")
    bold = Path(r"C:\Windows\Fonts\Dengb.ttf")
    if not regular.exists() or not bold.exists():
        regular = Path(r"C:\Windows\Fonts\simhei.ttf")
        bold = Path(r"C:\Windows\Fonts\simhei.ttf")
    pdfmetrics.registerFont(TTFont("ReportCN", str(regular)))
    pdfmetrics.registerFont(TTFont("ReportCN-Bold", str(bold)))


def clean(value: Any) -> Any:
    if value is None:
        return ""
    return value


def to_float(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def pct(value: Any) -> str:
    if value is None or value == "":
        return ""
    return f"{to_float(value) * 100:.2f}%"


def num(value: Any) -> str:
    if value is None or value == "":
        return ""
    try:
        return f"{float(value):,.0f}"
    except (TypeError, ValueError):
        return str(value)


def money(value: Any) -> str:
    if value is None or value == "":
        return ""
    try:
        return f"${float(value):,.2f}"
    except (TypeError, ValueError):
        return ""


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text in {"", "-", "(空)", "（空）", "None", "nan"}:
        return ""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def paragraph(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(text, style)


def first_existing_sheet(wb, names: Iterable[str]):
    for name in names:
        if name in wb.sheetnames:
            return wb[name]
    raise KeyError(f"Missing sheets: {list(names)}")


def read_core_metrics(wb) -> Dict[str, Dict[str, Any]]:
    ws = wb["00_核心对比"]
    metrics: Dict[str, Dict[str, Any]] = {}
    metric_aliases = [
        "total_transactions",
        "illion_coverage",
        "finv_coverage",
        "joint_agreement",
        "joint_difference",
        "adjusted_agreement",
        "adjusted_difference",
        "coverage_overlap",
        "coverage_nonoverlap",
        "illion_directional_difference",
        "finv_directional_difference",
        "difference_total",
        "illion_only_total",
        "finv_only_total",
        "both_empty_total",
    ]
    for row_index, row in enumerate(ws.iter_rows(min_row=6, max_row=20, min_col=1, max_col=5, values_only=True)):
        label = str(row[0] or "").strip()
        if label:
            value = {"result": row[1], "numerator": row[2], "denominator": row[3], "note": row[4]}
            metrics[label] = value
            metrics[metric_aliases[row_index]] = value
    return metrics


def read_categories(wb) -> List[Dict[str, Any]]:
    ws = wb["00_核心对比"]
    headers = [str(v or "") for v in next(ws.iter_rows(min_row=23, max_row=23, min_col=1, max_col=16, values_only=True))]
    items: List[Dict[str, Any]] = []
    for row in ws.iter_rows(min_row=24, max_row=59, min_col=1, max_col=16, values_only=True):
        if not row[0]:
            continue
        item = {headers[i]: row[i] for i in range(len(headers))}
        item["category"] = str(row[0])
        item["group"] = GROUP_OF.get(item["category"], "其他")
        item["illion_only_count"] = to_float(row[9])
        item["finv_only_count"] = to_float(row[10])
        item["side_diff_count"] = item["illion_only_count"] + item["finv_only_count"]
        item["illion_coverage"] = row[5]
        item["finv_coverage"] = row[6]
        item["union_count"] = row[7]
        item["intersection_count"] = row[8]
        item["intersection_share"] = row[11]
        item["illion_only_share"] = row[12]
        item["finv_only_share"] = row[13]
        item["priority"] = row[2]
        items.append(item)
    return items


def read_flows(wb) -> List[Dict[str, Any]]:
    ws = wb["01_差异诊断地图"]
    header_values = list(next(ws.iter_rows(min_row=5, max_row=5, min_col=1, max_col=15, values_only=True)))
    headers = [str(v or "") for v in header_values]
    flows: List[Dict[str, Any]] = []
    for row in ws.iter_rows(min_row=6, max_row=25, min_col=1, max_col=15, values_only=True):
        if not isinstance(row[0], (int, float)):
            continue
        flow = {headers[i]: row[i] for i in range(len(headers))}
        flow.update({
            "rank": row[0],
            "priority": row[1],
            "illion_category": row[3],
            "finv_category": row[4],
            "diff_type": row[5],
            "count": row[6],
            "diff_amount": row[12],
            "direction": row[13],
            "suggestion": row[14],
        })
        flows.append(flow)
    return flows


def read_detail(wb) -> List[Dict[str, Any]]:
    ws = wb["03_排查明细"]
    headers = [str(v or "") for v in next(ws.iter_rows(min_row=3, max_row=3, min_col=1, max_col=26, values_only=True))]
    rows: List[Dict[str, Any]] = []
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=26, values_only=True):
        if not any(v not in (None, "") for v in row):
            continue
        detail = {headers[i]: row[i] for i in range(len(headers))}
        detail.update({
            "priority": row[0],
            "diff_type": row[1],
            "diff_flow": row[2],
            "amount": row[6],
            "dr_cr": row[7],
            "text": row[8],
            "illion_category": row[11],
            "finv_category": row[12],
            "user_id": row[16],
            "application_id": row[17],
        })
        rows.append(detail)
    return rows


def detail_category_stats(details: List[Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    stats: Dict[str, Dict[str, Any]] = defaultdict(lambda: {
        "illion_diff": 0,
        "finv_diff": 0,
        "mismatch": 0,
        "illion_only": 0,
        "finv_only": 0,
        "amount": 0.0,
        "users": set(),
        "applications": set(),
    })
    for row in details:
        illion = str(row.get("illion Category") or "").strip()
        finv = str(row.get("finv Category") or "").strip()
        diff_type = str(row.get("排查类型") or "").strip()
        amount = to_float(row.get("amount"))
        users = row.get("user_id")
        applications = row.get("application_id")
        for category, side in ((illion, "illion_diff"), (finv, "finv_diff")):
            if category and category != "-":
                stats[category][side] += 1
                stats[category]["amount"] += amount / 2.0
                if users not in (None, ""):
                    stats[category]["users"].add(users)
                if applications not in (None, ""):
                    stats[category]["applications"].add(applications)
        if diff_type == "分类边界冲突":
            if illion and illion != "-":
                stats[illion]["mismatch"] += 1
            if finv and finv != "-":
                stats[finv]["mismatch"] += 1
        elif diff_type == "仅illion有分类":
            if illion and illion != "-":
                stats[illion]["illion_only"] += 1
        elif diff_type == "仅finv有分类":
            if finv and finv != "-":
                stats[finv]["finv_only"] += 1
    return stats


def category_conclusion(item: Dict[str, Any]) -> str:
    category = safe_text(item["category"])
    illion_cov = to_float(item.get("illion覆盖率"))
    finv_cov = to_float(item.get("finv覆盖率"))
    inter_share = to_float(item.get("交集占比（并集）"))
    illion_only = to_float(item.get("illion独有占比（并集）"))
    finv_only = to_float(item.get("finv独有占比（并集）"))
    union = to_float(item.get("并集数量"))
    priority = safe_text(item.get("建议优先级"))
    if union < 50:
        volume_note = "并集规模较小，比例指标需结合样本量谨慎解读。"
    elif union >= 1000:
        volume_note = "并集规模较大，对总体差异具有较高影响。"
    else:
        volume_note = "并集规模处于中等水平。"
    if inter_share >= 0.9:
        stability = "双方交集占比较高，分类表现相对稳定。"
    elif inter_share >= 0.7:
        stability = "双方存在较好的共同覆盖，但仍有一定单边差异。"
    else:
        stability = "双方共同覆盖偏低，应重点检查分类边界、知识库或漏识别问题。"
    if finv_only > illion_only + 0.05:
        direction = "finv 独有占比明显高于 Illion 独有占比，体现 finv 的扩展识别，同时需要核验新增分类的合理性。"
    elif illion_only > finv_only + 0.05:
        direction = "Illion 独有占比高于 finv 独有占比，提示 finv 可能存在漏识别或归类迁移。"
    else:
        direction = "两侧独有占比较为接近，差异更可能来自类别边界或规则口径。"
    return f"{category}：Illion 覆盖率为 {pct(illion_cov)}，finv 覆盖率为 {pct(finv_cov)}；并集 {num(union)} 笔，交集占比 {pct(inter_share)}。{stability}{direction}{volume_note} 建议优先级为 {priority or '-'}。"


def flow_rows_for_category(flows: List[Dict[str, Any]], category: str, limit: int = 3) -> List[Dict[str, Any]]:
    matched = []
    for flow in flows:
        source = str(flow.get("illion Category") or "")
        target = str(flow.get("finv Category") or "")
        if category in source or category in target:
            matched.append(flow)
    return matched[:limit]


def table(data: List[List[Any]], widths: List[float], header_rows: int = 1, font_size: float = 7.5, row_bgs: bool = True) -> Table:
    header_style = ParagraphStyle(
        "_table_header",
        fontName="ReportCN-Bold",
        fontSize=font_size,
        leading=font_size + 2,
        textColor=colors.white,
        wordWrap="CJK",
    )
    cell_style = ParagraphStyle(
        "_table_cell",
        fontName="ReportCN",
        fontSize=font_size,
        leading=font_size + 2,
        textColor=colors.HexColor("#222222"),
        wordWrap="CJK",
    )
    normalized: List[List[Any]] = []
    for row_index, row in enumerate(data):
        row_cells = []
        for x in row:
            if isinstance(x, Paragraph):
                row_cells.append(x)
            else:
                row_cells.append(Paragraph(safe_text(x), header_style if row_index < header_rows else cell_style))
        normalized.append(row_cells)
    t = Table(normalized, colWidths=widths, repeatRows=header_rows, hAlign="LEFT")
    commands = [
        ("BACKGROUND", (0, 0), (-1, header_rows - 1), BLUE),
        ("TEXTCOLOR", (0, 0), (-1, header_rows - 1), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#A7B4C1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if row_bgs:
        for idx in range(header_rows, len(data)):
            if (idx - header_rows) % 2 == 1:
                commands.append(("BACKGROUND", (0, idx), (-1, idx), BLUE_PALE))
    t.setStyle(TableStyle(commands))
    return t


class ReportDocTemplate(BaseDocTemplate):
    def __init__(self, filename: str, **kwargs):
        super().__init__(filename, **kwargs)
        frame = Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id="normal")
        self.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=draw_page)])


def draw_page(canvas, doc) -> None:
    canvas.saveState()
    width, height = A4
    canvas.setStrokeColor(colors.HexColor("#D0D7DE"))
    canvas.setLineWidth(0.4)
    canvas.line(doc.leftMargin, height - 16 * mm, width - doc.rightMargin, height - 16 * mm)
    canvas.setFont("ReportCN", 7.5)
    canvas.setFillColor(GRAY)
    canvas.drawString(doc.leftMargin, 10 * mm, "BS-CAT Full Category Performance Report | 数据来源: category_difference_report(2).xlsx")
    canvas.drawRightString(width - doc.rightMargin, 10 * mm, f"第 {doc.page} 页")
    canvas.restoreState()


def build_styles() -> Dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("title", parent=base["Title"], fontName="ReportCN-Bold", fontSize=23, leading=30, alignment=TA_CENTER, textColor=BLUE, spaceAfter=8),
        "subtitle": ParagraphStyle("subtitle", parent=base["Normal"], fontName="ReportCN", fontSize=10, leading=15, alignment=TA_CENTER, textColor=GRAY, spaceAfter=18),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontName="ReportCN-Bold", fontSize=16, leading=22, textColor=BLUE, spaceBefore=10, spaceAfter=8),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontName="ReportCN-Bold", fontSize=12, leading=18, textColor=colors.HexColor("#2873B2"), spaceBefore=8, spaceAfter=5),
        "h3": ParagraphStyle("h3", parent=base["Heading3"], fontName="ReportCN-Bold", fontSize=10.5, leading=16, textColor=BLUE, spaceBefore=6, spaceAfter=4),
        "body": ParagraphStyle("body", parent=base["BodyText"], fontName="ReportCN", fontSize=9.2, leading=15, textColor=colors.HexColor("#222222"), spaceAfter=6, wordWrap="CJK"),
        "small": ParagraphStyle("small", parent=base["BodyText"], fontName="ReportCN", fontSize=7.6, leading=11, textColor=GRAY, spaceAfter=4, wordWrap="CJK"),
        "callout": ParagraphStyle("callout", parent=base["BodyText"], fontName="ReportCN", fontSize=9, leading=15, textColor=colors.HexColor("#2F2F2F"), leftIndent=10, rightIndent=8, spaceAfter=5, wordWrap="CJK"),
        "table": ParagraphStyle("table", parent=base["BodyText"], fontName="ReportCN", fontSize=7.2, leading=9.5, wordWrap="CJK"),
        "table_bold": ParagraphStyle("table_bold", parent=base["BodyText"], fontName="ReportCN-Bold", fontSize=7.2, leading=9.5, wordWrap="CJK"),
        "center": ParagraphStyle("center", parent=base["BodyText"], fontName="ReportCN", fontSize=8, leading=11, alignment=TA_CENTER),
    }


def add_section_header(story: List[Any], title: str, styles: Dict[str, ParagraphStyle]) -> None:
    story.append(Spacer(1, 4))
    story.append(paragraph(title, styles["h1"]))
    rule = Table([[""]], colWidths=[170 * mm], rowHeights=[1.2])
    rule.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), BLUE), ("LINEBELOW", (0, 0), (-1, -1), 0, BLUE)]))
    story.append(rule)
    story.append(Spacer(1, 5))


def add_callout(story: List[Any], text: str, styles: Dict[str, ParagraphStyle], background=BLUE_PALE, edge=BLUE) -> None:
    box = Table([[paragraph(text, styles["callout"])]], colWidths=[170 * mm])
    box.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), background),
        ("LINEBEFORE", (0, 0), (0, -1), 4, edge),
        ("BOX", (0, 0), (-1, -1), 0.25, background),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(box)
    story.append(Spacer(1, 5))


def category_metric_table(item: Dict[str, Any], styles: Dict[str, ParagraphStyle]) -> Table:
    data = [
        ["类别", "Illion覆盖率", "finv覆盖率", "并集数量", "交集数量", "交集占比（并集）", "Illion独有占比", "finv独有占比"],
        [
            safe_text(item["category"]),
            pct(item.get("illion覆盖率")),
            pct(item.get("finv覆盖率")),
            num(item.get("并集数量")),
            num(item.get("交集数量")),
            pct(item.get("交集占比（并集）")),
            pct(item.get("illion独有占比（并集）")),
            pct(item.get("finv独有占比（并集）")),
        ],
    ]
    widths = [28 * mm, 20 * mm, 20 * mm, 16 * mm, 16 * mm, 22 * mm, 22 * mm, 22 * mm]
    return table(data, widths, font_size=6.9)


def add_category_block(story: List[Any], item: Dict[str, Any], flows: List[Dict[str, Any]], detail_stats: Dict[str, Dict[str, Any]], styles: Dict[str, ParagraphStyle], detailed: bool = False) -> None:
    category = item["category"]
    heading = f"{category}  |  {item['group']}"
    block: List[Any] = [paragraph(heading, styles["h3"]), category_metric_table(item, styles)]
    block.append(Spacer(1, 3))
    block.append(paragraph(category_conclusion(item), styles["body"]))
    extra = detail_stats.get(category, {})
    if detailed:
        block.append(paragraph(
            f"差异侧统计：Illion 侧差异记录 {num(extra.get('illion_diff'))} 笔，finv 侧差异记录 {num(extra.get('finv_diff'))} 笔；"
            f"其中双方分类不一致 {num(extra.get('mismatch'))} 笔。该侧统计用于定位差异集中区域，不与总体 6,525 笔差异直接相加。",
            styles["small"],
        ))
        flow_matches = flow_rows_for_category(flows, category, limit=3)
        if flow_matches:
            rows = [["排名", "Illion类别", "finv类别", "类型", "数量", "差异金额"]]
            for idx, flow in enumerate(flow_matches, start=1):
                rows.append([
                    idx,
                    safe_text(flow.get("illion Category")),
                    safe_text(flow.get("finv Category")),
                    safe_text(flow.get("差异类型")),
                    num(flow.get("数量")),
                    money(flow.get("差异金额")),
                ])
            block.append(table(rows, [12 * mm, 34 * mm, 34 * mm, 28 * mm, 18 * mm, 25 * mm], font_size=7.0))
    else:
        block.append(paragraph("该类别按统一七项指标完成分析；详细差异流向见附录类别评分表。", styles["small"]))
    story.append(KeepTogether(block))
    story.append(Spacer(1, 5))


def build_report(input_path: Path, output_path: Path) -> Dict[str, Any]:
    register_fonts()
    wb = load_workbook(input_path, data_only=True, read_only=True)
    metrics = read_core_metrics(wb)
    categories = read_categories(wb)
    flows = read_flows(wb)
    details = read_detail(wb)
    detail_stats = detail_category_stats(details)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = ReportDocTemplate(
        str(output_path),
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=23 * mm,
        bottomMargin=17 * mm,
        title="BS-CAT 全分类类别表现分析报告",
        author="Codex",
    )
    styles = build_styles()
    story: List[Any] = []

    total_transactions = metrics.get("总交易数", {}).get("result", 0)
    illion_cov = metrics.get("illion Category 覆盖率", {}).get("result", 0)
    finv_cov = metrics.get("finv Category 覆盖率", {}).get("result", 0)
    joint_agreement = metrics.get("双方非空时一致率", {}).get("result", 0)
    adjusted_agreement = metrics.get("覆盖调整后一致率", {}).get("result", 0)
    direction_illion = metrics.get("相对illion的方向性差异率", {}).get("result", 0)
    direction_finv = metrics.get("相对finv的方向性差异率", {}).get("result", 0)
    diff_total = metrics.get("Category 差异总数", {}).get("result", 0)
    illion_only = metrics.get("仅 illion 有 Category", {}).get("result", 0)
    finv_only = metrics.get("仅 finv 有 Category", {}).get("result", 0)
    both_empty = metrics.get("双方均为空", {}).get("result", 0)
    mismatch = metrics.get("双方非空时差异率", {}).get("numerator", 0)

    # Cover page
    story.append(Spacer(1, 18 * mm))
    story.append(paragraph("BS-CAT 全分类类别表现分析报告", styles["title"]))
    story.append(paragraph("Income, Expense, Transfer and Liability Category Performance", styles["subtitle"]))
    story.append(Spacer(1, 8 * mm))
    cover_table = table([
        ["样本范围", "分类范围", "数据侧", "报告日期"],
        [f"{num(total_transactions)} 笔交易", "36 个类别", "illion vs finv / BS-CAT", "2026-08-19"],
    ], [42 * mm, 42 * mm, 42 * mm, 42 * mm], font_size=8.5)
    story.append(cover_table)
    story.append(Spacer(1, 16 * mm))
    add_callout(story, "本报告基于 category_difference_report(2).xlsx 自动生成。报告将收入、支出、转账和负债类别统一纳入全分类分析；负债部分保留类别指标摘要，不重复展开已有 Liability 专项报告。", styles)
    story.append(Spacer(1, 24 * mm))
    story.append(paragraph("数据来源", styles["h2"]))
    story.append(paragraph("Excel 底稿：00_核心对比、01_差异诊断地图、03_排查明细、04_模型监控。类别指标以 00_核心对比中的 36 个类别行作为主口径。", styles["body"]))
    story.append(PageBreak())

    # 1 Summary
    add_section_header(story, "1. Summary", styles)
    story.append(paragraph(
        f"本次评估包含 {num(total_transactions)} 笔交易和 36 个类别，覆盖收入、支出、负债三大业务类别，并将 External Transfers 与 Internal Transfer 作为独立的转账专项进行分析。",
        styles["body"],
    ))
    joint_nonempty = metrics.get("双方非空时一致率", {}).get("denominator", 0)
    add_callout(
        story,
        f"总体结论：在 {num(total_transactions)} 笔交易中，illion 分类覆盖率为 {pct(illion_cov)}，finv 分类覆盖率为 {pct(finv_cov)}，finv 覆盖率高于 Illion {pct(to_float(finv_cov) - to_float(illion_cov))}。在双方均有非空分类的 {num(joint_nonempty)} 笔交易中，分类一致率为 {pct(joint_agreement)}。因此，finv 的主要表现特征是覆盖范围更广，同时双方在已完成分类的交易上保持较高的一致性。",
        styles,
    )

    story.append(paragraph("1.1 差异结构总览", styles["h2"]))
    diff_table = table([
        ["差异类型", "数量", "占差异总数", "分析含义"],
        ["双方分类不一致", num(mismatch), pct(to_float(mismatch) / to_float(diff_total) if diff_total else 0), "两侧均有分类，但类别标签不同"],
        ["仅 illion 有值", num(illion_only), pct(to_float(illion_only) / to_float(diff_total) if diff_total else 0), "finv 侧缺少分类"],
        ["仅 finv 有值", num(finv_only), pct(to_float(finv_only) / to_float(diff_total) if diff_total else 0), "finv 新增识别或 Illion 漏识别"],
        ["总差异数", num(diff_total), "100.00%", "分类不一致与单边缺失的合计"],
        ["双方均为空", num(both_empty), "不纳入", "不进入差异流向和类别差异分析"],
    ], [38 * mm, 22 * mm, 25 * mm, 85 * mm], font_size=7.8)
    story.append(diff_table)
    add_callout(story, f"仅 finv 有值共 {num(finv_only)} 笔，占差异总数 {pct(to_float(finv_only) / to_float(diff_total) if diff_total else 0)}，是最大差异来源。差异集中区域主要包括转账类、高频消费类和贷款类；其中 Rent 作为支出类别单独进行重点分析。", styles, background=YELLOW_PALE, edge=ORANGE)

    # 2 scope
    add_section_header(story, "2. Category Grouping and Analysis Scope", styles)
    story.append(paragraph("本报告使用三大业务类别加一个转账专项的分析口径。转账类不归入收入或支出，不参与收入合计和支出合计，但纳入总体覆盖率、差异率和分类迁移影响分析。每个类别均按照 Illion 覆盖率、finv 覆盖率、并集数量、交集数量、交集占比（并集）、Illion 独有占比和 finv 独有占比七项核心指标进行分析。", styles["body"]))
    grouping_rows = [["分析组", "类别数", "类别范围", "报告处理"]]
    handling = {
        "收入类": "完整分析",
        "支出类": "完整分析，Rent 重点展开",
        "负债类": "保留七项指标，压缩为摘要",
        "转账专项": "独立分析，不并入收入或支出",
    }
    for group, categories_in_group in GROUPS.items():
        grouping_rows.append([group, len(categories_in_group), ", ".join(categories_in_group), handling[group]])
    story.append(table(grouping_rows, [24 * mm, 14 * mm, 94 * mm, 38 * mm], font_size=6.9))
    story.append(PageBreak())

    # 3 overall performance
    add_section_header(story, "3. Overall Category Performance", styles)
    story.append(paragraph("36 个类别的主指标来自底稿中的逐 Category 对比表。交集占比用于衡量该类别的共同覆盖程度；两侧独有占比用于识别新增识别、漏识别和分类边界问题。", styles["body"]))
    top_union = sorted(categories, key=lambda x: to_float(x.get("并集数量")), reverse=True)[:10]
    top_diff = sorted(categories, key=lambda x: to_float(x.get("side_diff_count")), reverse=True)[:10]
    story.append(paragraph("3.1 并集规模最大的类别", styles["h2"]))
    union_rows = [["排名", "类别", "分析组", "并集数量", "交集数量", "交集占比", "finv独有占比"]]
    for idx, item in enumerate(top_union, 1):
        union_rows.append([idx, item["category"], item["group"], num(item.get("并集数量")), num(item.get("交集数量")), pct(item.get("交集占比（并集）")), pct(item.get("finv独有占比（并集）"))])
    story.append(table(union_rows, [11 * mm, 36 * mm, 23 * mm, 22 * mm, 22 * mm, 25 * mm, 25 * mm], font_size=7.5))
    story.append(Spacer(1, 7))
    story.append(paragraph("3.2 类别侧差异集中度", styles["h2"]))
    story.append(paragraph("以下按 Illion 独有数量与 finv 独有数量之和进行类别侧差异排名。该排名用于定位差异集中类别，不能将各类别差异数量再次相加作为总体差异总数。", styles["small"]))
    diff_rows = [["排名", "类别", "分析组", "类别并集", "交集占比", "Illion独有数", "finv独有数", "类别侧差异数"]]
    for idx, item in enumerate(top_diff, 1):
        diff_rows.append([idx, item["category"], item["group"], num(item.get("并集数量")), pct(item.get("交集占比（并集）")), num(item.get("illion_only_count")), num(item.get("finv_only_count")), num(item.get("side_diff_count"))])
    story.append(table(diff_rows, [10 * mm, 31 * mm, 22 * mm, 18 * mm, 22 * mm, 21 * mm, 21 * mm, 24 * mm], font_size=7.1))
    story.append(Spacer(1, 7))
    story.append(PageBreak())
    story.append(paragraph("3.3 Top 差异流向", styles["h2"]))
    flow_rows = [["排名", "Illion类别", "finv类别", "差异类型", "数量", "差异金额", "优先级"]]
    for idx, flow in enumerate(flows[:20], 1):
        flow_rows.append([idx, safe_text(flow.get("illion Category")), safe_text(flow.get("finv Category")), safe_text(flow.get("差异类型")), num(flow.get("数量")), money(flow.get("差异金额")), safe_text(flow.get("建议优先级"))])
    story.append(table(flow_rows, [10 * mm, 30 * mm, 30 * mm, 28 * mm, 16 * mm, 25 * mm, 18 * mm], font_size=6.9))
    story.append(PageBreak())

    # 4 difference decomposition
    add_section_header(story, "4. Difference Decomposition", styles)
    story.append(paragraph("差异分析分为两个层次：第一层是交易样本层面的总体差异结构；第二层是类别侧的差异集中度。总体差异采用互斥口径，类别侧排名用于定位重点类别，不与总体差异总数直接相加。", styles["body"]))
    story.append(paragraph("4.1 总体差异结构", styles["h2"]))
    story.append(diff_table)
    story.append(Spacer(1, 7))
    story.append(paragraph("4.2 差异集中类别的拆解", styles["h2"]))
    story.append(paragraph("类别层面同时观察两侧覆盖率、并集规模、交集规模和两侧独有占比。高并集且低交集占比的类别是优先排查对象；高 finv 独有占比类别需要进一步验证新增识别是否合理；高 Illion 独有占比类别需要排查 finv 漏识别。", styles["body"]))
    story.append(table(diff_rows, [10 * mm, 31 * mm, 22 * mm, 18 * mm, 22 * mm, 21 * mm, 21 * mm, 24 * mm], font_size=7.1))
    story.append(Spacer(1, 7))
    story.append(paragraph("4.3 差异流向的业务解释", styles["h2"]))
    story.append(paragraph("重点解释转账与收入的边界、高频消费类别之间的边界，以及 finv 独有分类的覆盖扩展。Rent 在支出模块中单列，不因类别侧差异排名低于转账类而降低分析优先级。", styles["body"]))
    story.append(PageBreak())

    # 5 income
    add_section_header(story, "5. Income Category Analysis", styles)
    story.append(paragraph("收入类包括 Wages、Centrelink 和 All Other Credits，不包含 External Transfers 与 Internal Transfer。每个类别使用七项核心指标；收入与转账之间的分类迁移在第 7 章单独说明。", styles["body"]))
    for item in categories:
        if item["group"] == "收入类":
            add_category_block(story, item, flows, detail_stats, styles, detailed=True)
    story.append(PageBreak())

    # 6 expense
    add_section_header(story, "6. Expense Category Analysis", styles)
    story.append(paragraph("支出类共 23 个类别。类别层面统一展示七项核心指标；Groceries、Dining Out、Retail、Automotive、Gambling、Utilities 和 Rent 作为高频或重点边界类别增加差异流向说明。", styles["body"]))
    for item in categories:
        if item["group"] == "支出类":
            add_category_block(story, item, flows, detail_stats, styles, detailed=item["category"] in FOCUS_CATEGORIES or item["category"] == "Utilities")
    story.append(PageBreak())

    # Rent focus
    rent = next((item for item in categories if item["category"] == "Rent"), None)
    if rent:
        add_section_header(story, "6.1 Rent Deep Dive", styles)
        story.append(paragraph("Rent 是支出模块的重点类别。除七项核心指标外，本节关注 Rent 与转账、Utilities 和 Home Improvement 的边界，重点判断 finv 独有识别是合理扩展还是潜在误归类。", styles["body"]))
        add_category_block(story, rent, flows, detail_stats, styles, detailed=True)
        add_callout(story, "Rent 的类别结论应同时结合覆盖率、交集占比和差异流向，不应只根据交易数量或类别差异率判断模型表现。", styles, background=YELLOW_PALE, edge=ORANGE)
        story.append(PageBreak())

    # 7 transfers
    add_section_header(story, "7. Transfer Special Analysis", styles)
    story.append(paragraph("External Transfers 和 Internal Transfer 不归入收入或支出。两类转账按 credit / debit 方向单独观察，并分析其与 Wages、All Other Credits 和 Rent 的分类边界。", styles["body"]))
    for item in categories:
        if item["group"] == "转账专项":
            add_category_block(story, item, flows, detail_stats, styles, detailed=True)
    add_callout(story, "转账类的报告处理原则：保留在总体覆盖率、差异率和分类迁移分析中，但不计入收入合计、支出合计或净收入计算。", styles)
    story.append(PageBreak())

    # 8 liability summary
    add_section_header(story, "8. Liability Category Summary", styles)
    story.append(paragraph("负债类在本报告中保留 8 个类别的统一七项指标，用于完整覆盖 36 个类别的总体表现；贷款生命周期、Counterparty matching、Dishonours 和 Unknown Loans 根因不在本报告重复展开。", styles["body"]))
    for item in categories:
        if item["group"] == "负债类":
            add_category_block(story, item, flows, detail_stats, styles, detailed=False)
    add_callout(story, "Liability Module 的详细结论应以已有专项报告为准。本报告仅保留负债类别在全分类差异结构中的指标表现和类别侧定位。", styles, background=YELLOW_PALE, edge=ORANGE)
    story.append(PageBreak())

    # 9 action plan
    add_section_header(story, "9. Cross-Category Action Plan", styles)
    action_rows = [
        ["优先方向", "建议动作", "对应类别"],
        ["验证 finv 独有识别", "抽样核验新增分类是否合理，区分覆盖扩展与误归类", "External Transfers、All Other Credits、Retail、Dining Out、Unknown Loans"],
        ["排查 Illion 独有识别", "检查 finv 是否存在漏识别、文本清洗或知识库缺口", "External Transfers、Groceries、Non SACC Loans、Gambling"],
        ["优化转账边界", "拆分 credit/debit，明确中性转移不进入收入或支出合计", "External Transfers、Internal Transfer"],
        ["优化 Rent 边界", "重点复核 Rent 与转账、Utilities、Home Improvement 的规则边界", "Rent、Utilities、Home Improvement"],
        ["优化高频消费边界", "针对高并集、低交集占比类别建立样本集和规则回归集", "Groceries、Dining Out、Retail、Automotive、Gambling"],
        ["引用专项结果", "不在本报告重复做贷款生命周期和 Dishonours 深度分析", "全部负债类别"],
    ]
    story.append(table(action_rows, [29 * mm, 89 * mm, 52 * mm], font_size=7.4))
    story.append(Spacer(1, 8))
    story.append(paragraph("后续迭代建议：每次模型或知识库更新后，重新计算 36 个类别的七项核心指标，并重点监控交集占比、finv 独有占比、Illion 独有占比以及 Rent 和转账类的差异流向。", styles["body"]))
    story.append(PageBreak())

    # 10 appendix full scorecard
    add_section_header(story, "10. Appendix: 36-Category Scorecard", styles)
    story.append(paragraph("以下为 36 个类别的完整七项指标。该表是本报告的核心可审计明细，所有类别均使用相同口径。注：类别侧指标用于解释该类别的共同覆盖和单边差异；类别之间的交叉错分会同时出现在源类别和目标类别的类别侧视角中，因此类别侧差异排名不能直接加总为总体差异数。", styles["body"]))
    score_rows = [["类别", "分析组", "Illion覆盖率", "finv覆盖率", "并集", "交集", "交集占比", "Illion独有占比", "finv独有占比"]]
    for item in categories:
        score_rows.append([
            item["category"], item["group"], pct(item.get("illion覆盖率")), pct(item.get("finv覆盖率")),
            num(item.get("并集数量")), num(item.get("交集数量")), pct(item.get("交集占比（并集")),
            pct(item.get("illion独有占比（并集")), pct(item.get("finv独有占比（并集")),
        ])
    # The source headers contain closing Chinese brackets; normalize if a key lookup above misses.
    for row_idx, item in enumerate(categories, start=1):
        score_rows[row_idx] = [
            item["category"], item["group"], pct(item.get("illion覆盖率")), pct(item.get("finv覆盖率")),
            num(item.get("并集数量")), num(item.get("交集数量")), pct(item.get("交集占比（并集）")),
            pct(item.get("illion独有占比（并集）")), pct(item.get("finv独有占比（并集）")),
        ]
    story.append(table(score_rows, [26 * mm, 20 * mm, 18 * mm, 18 * mm, 14 * mm, 14 * mm, 19 * mm, 20 * mm, 20 * mm], font_size=6.2, row_bgs=True))

    doc.build(story)
    return {
        "output": str(output_path),
        "categories": len(categories),
        "details": len(details),
        "flows": len(flows),
        "metrics": metrics,
    }


# The rebuilt report uses stable ASCII aliases internally and keeps the Chinese labels
# in one place, so the PDF logic is independent of the workbook's display-language headers.
REPORT_GROUPS = {
    "income": ["Wages", "Centrelink", "All Other Credits"],
    "expense": [
        "Information", "Donations", "Education", "Home Improvement", "Insurance",
        "Subscription TV", "Pet Care", "Entertainment", "Retail", "Utilities",
        "Gyms and other memberships", "Personal Care", "Rent", "Health", "Travel",
        "Groceries", "Automotive", "Gambling", "Department Stores", "Telecommunications",
        "Dining Out", "Fees", "Transport",
    ],
    "transfer": ["External Transfers", "Internal Transfer"],
    "liability": [
        "SACC Loans", "Non SACC Loans", "Credit Card Repayments", "Debt Consolidation",
        "Debt Collection", "Dishonours", "Overdrawn", "Unknown Loans",
    ],
}
REPORT_GROUP_LABELS = {
    "income": "收入类",
    "expense": "支出类",
    "transfer": "转账类",
    "liability": "负债类",
}
REPORT_GROUP_ORDER = ["income", "expense", "transfer", "liability"]
REPORT_GROUP_OF = {category: group for group, values in REPORT_GROUPS.items() for category in values}
REPORT_FOCUS = {"Wages", "All Other Credits", "Rent", "Groceries", "Dining Out", "Retail", "Automotive", "Gambling", "Utilities", "External Transfers", "Internal Transfer"}


def report_group(category: Any) -> str:
    value = str(category or "").strip()
    return REPORT_GROUP_OF.get(value, "")


def is_missing_category(value: Any) -> bool:
    return str(value or "").strip() in {"", "-", "(空)", "（空）", "None"}


def build_block_analysis(details: List[Dict[str, Any]]) -> Dict[str, Any]:
    empty_illion = "Illion为空"
    empty_finv = "finv为空"
    rows = REPORT_GROUP_ORDER + [empty_illion]
    cols = REPORT_GROUP_ORDER + [empty_finv]
    matrix = {(row, col): 0 for row in rows for col in cols}
    for detail in details:
        illion = str(detail.get("illion_category") or "").strip()
        finv = str(detail.get("finv_category") or "").strip()
        source = report_group(illion) if not is_missing_category(illion) else empty_illion
        target = report_group(finv) if not is_missing_category(finv) else empty_finv
        if (source, target) in matrix:
            matrix[(source, target)] += 1
    summary = {}
    for group in REPORT_GROUP_ORDER:
        summary[group] = {
            "illion_side": sum(matrix[(group, col)] for col in cols),
            "finv_side": sum(matrix[(row, group)] for row in rows),
            "within": matrix[(group, group)],
            "source_cross": sum(matrix[(group, other)] for other in REPORT_GROUP_ORDER if other != group),
            "target_cross": sum(matrix[(other, group)] for other in REPORT_GROUP_ORDER if other != group),
            "illion_only": matrix[(group, empty_finv)],
            "finv_only": matrix[(empty_illion, group)],
        }
    return {"rows": rows, "cols": cols, "matrix": matrix, "summary": summary}


def category_narrative_v2(item: Dict[str, Any]) -> str:
    category = safe_text(item.get("category"))
    illion_coverage = pct(item.get("illion_coverage"))
    finv_coverage = pct(item.get("finv_coverage"))
    union_count = num(item.get("union_count"))
    intersection_count = num(item.get("intersection_count"))
    intersection_share = pct(item.get("intersection_share"))
    illion_only_share = pct(item.get("illion_only_share"))
    finv_only_share = pct(item.get("finv_only_share"))
    priority = safe_text(item.get("priority"))
    if item.get("union_count") in (None, ""):
        scale_text = "并集规模：。"
    elif to_float(item.get("union_count")) >= 1000:
        scale_text = "并集规模较大，对整体差异影响较高。"
    elif to_float(item.get("union_count")) < 50:
        scale_text = "并集规模较小，比例指标需结合样本量谨慎解读。"
    else:
        scale_text = "并集规模处于中等水平。"
    if item.get("intersection_share") in (None, ""):
        stability_text = "共同覆盖程度：。"
    elif to_float(item.get("intersection_share")) >= 0.9:
        stability_text = "交集占比较高，双方对该类别的共同识别较稳定。"
    elif to_float(item.get("intersection_share")) >= 0.7:
        stability_text = "双方存在较好的共同覆盖，但仍有一定单边差异。"
    else:
        stability_text = "交集占比较低，应重点排查类别边界、知识库或漏识别问题。"
    if item.get("finv_only_share") in (None, "") or item.get("illion_only_share") in (None, ""):
        direction_text = "两侧独有占比：。"
    elif to_float(item.get("finv_only_share")) > to_float(item.get("illion_only_share")) + 0.05:
        direction_text = "finv 独有占比更高，体现扩展识别，同时需要核验新增分类的合理性。"
    elif to_float(item.get("illion_only_share")) > to_float(item.get("finv_only_share")) + 0.05:
        direction_text = "Illion 独有占比更高，提示 finv 可能存在漏识别或归类迁移。"
    else:
        direction_text = "两侧独有占比较接近，差异更可能来自类别边界或规则口径。"
    return (
        f"{category}：Illion 覆盖率 {illion_coverage}，finv 覆盖率 {finv_coverage}；并集 {union_count} 笔，"
        f"交集 {intersection_count} 笔，交集占比（并集）{intersection_share}；Illion 独有占比 {illion_only_share}，"
        f"finv 独有占比 {finv_only_share}。{stability_text}{direction_text}{scale_text}建议优先级为 {priority}。"
    )


def category_metric_table_v2(item: Dict[str, Any]) -> Table:
    rows = [
        ["类别", "Illion覆盖率", "finv覆盖率", "并集数量", "交集数量", "交集占比（并集）", "Illion独有占比", "finv独有占比"],
        [
            safe_text(item.get("category")), pct(item.get("illion_coverage")), pct(item.get("finv_coverage")),
            num(item.get("union_count")), num(item.get("intersection_count")), pct(item.get("intersection_share")),
            pct(item.get("illion_only_share")), pct(item.get("finv_only_share")),
        ],
    ]
    return table(rows, [28 * mm, 20 * mm, 20 * mm, 16 * mm, 16 * mm, 22 * mm, 22 * mm, 22 * mm], font_size=6.8)


def category_flow_rows_v2(flows: List[Dict[str, Any]], category: str, limit: int = 3) -> List[Dict[str, Any]]:
    matched = []
    for flow in flows:
        if category in str(flow.get("illion_category") or "") or category in str(flow.get("finv_category") or ""):
            matched.append(flow)
    return matched[:limit]


def add_category_block_v2(story: List[Any], item: Dict[str, Any], flows: List[Dict[str, Any]], styles: Dict[str, ParagraphStyle], detailed: bool = False) -> None:
    category = str(item.get("category") or "")
    group_label = REPORT_GROUP_LABELS.get(report_group(category), "")
    block = [
        paragraph(f"{safe_text(category)} | {safe_text(group_label)}", styles["h3"]),
        category_metric_table_v2(item),
        Spacer(1, 3),
        paragraph(category_narrative_v2(item), styles["body"]),
    ]
    if detailed:
        flow_rows = category_flow_rows_v2(flows, category)
        if flow_rows:
            rows = [["排名", "Illion类别", "finv类别", "差异类型", "数量", "差异金额", "主要方向"]]
            for index, flow in enumerate(flow_rows, 1):
                rows.append([
                    index, safe_text(flow.get("illion_category")), safe_text(flow.get("finv_category")),
                    safe_text(flow.get("diff_type")), num(flow.get("count")), money(flow.get("diff_amount")),
                    safe_text(flow.get("direction")),
                ])
            block.append(table(rows, [10 * mm, 29 * mm, 29 * mm, 26 * mm, 16 * mm, 24 * mm, 30 * mm], font_size=6.8))
    story.append(KeepTogether(block))
    story.append(Spacer(1, 5))


def build_report_v2(input_path: Path, output_path: Path) -> Dict[str, Any]:
    register_fonts()
    wb = load_workbook(input_path, data_only=True, read_only=True)
    metrics = read_core_metrics(wb)
    categories = read_categories(wb)
    flows = read_flows(wb)
    details = read_detail(wb)
    block_analysis = build_block_analysis(details)
    category_count = len(categories)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = ReportDocTemplate(
        str(output_path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=23 * mm, bottomMargin=17 * mm,
        title="BS-CAT 全分类类别表现分析报告", author="Codex",
    )
    styles = build_styles()
    story: List[Any] = []

    total_transactions = metrics.get("total_transactions", {}).get("result")
    illion_coverage = metrics.get("illion_coverage", {}).get("result")
    finv_coverage = metrics.get("finv_coverage", {}).get("result")
    joint_agreement = metrics.get("joint_agreement", {}).get("result")
    joint_nonempty = metrics.get("joint_agreement", {}).get("denominator")
    diff_total = metrics.get("difference_total", {}).get("result")
    mismatch = metrics.get("joint_difference", {}).get("numerator")
    illion_only = metrics.get("illion_only_total", {}).get("result")
    finv_only = metrics.get("finv_only_total", {}).get("result")
    both_empty = metrics.get("both_empty_total", {}).get("result")

    # Cover
    story.append(Spacer(1, 18 * mm))
    story.append(paragraph("BS-CAT 全分类类别表现分析报告", styles["title"]))
    story.append(paragraph("全分类差异结构与业务板块类型细探", styles["subtitle"]))
    story.append(Spacer(1, 8 * mm))
    story.append(table([
        ["样本范围", "分类范围", "比较对象", "报告日期"],
        [f"{num(total_transactions)} 笔交易", "36 个类别", "illion vs finv / BS-CAT", "2026-08-19"],
    ], [42 * mm, 42 * mm, 42 * mm, 42 * mm], font_size=8.5))
    story.append(Spacer(1, 16 * mm))
    add_callout(story, "本报告按新的三部分逻辑生成：执行摘要、差异结构全景、分业务板块的类型细探。对底稿中没有提供的字段不作推断，相关位置留空。", styles)
    story.append(Spacer(1, 22 * mm))
    story.append(paragraph("数据来源", styles["h2"]))
    story.append(paragraph("Excel 底稿：00_核心对比、01_差异诊断地图、03_排查明细、04_模型监控。", styles["body"]))
    story.append(PageBreak())

    # Part 1: Execution summary
    add_section_header(story, "1. 执行摘要", styles)
    story.append(paragraph(
        f"本次评估覆盖 {num(total_transactions)} 笔交易和 {category_count} 个类别，涉及收入、支出、负债三大业务板块，并将 External Transfers 与 Internal Transfer 作为独立的转账板块进行分析。Illion 分类覆盖率为 {pct(illion_coverage)}，finv 分类覆盖率为 {pct(finv_coverage)}，finv 的覆盖范围高于 Illion；在双方均有非空分类的 {num(joint_nonempty)} 笔交易中，分类一致率达到 {pct(joint_agreement)}，说明双方在已完成分类的交易上具有较高的一致性。总体来看，finv 的主要特征是覆盖更广，但新增覆盖也带来了更多差异；当前 {num(diff_total)} 笔差异中，仅 finv 有值占 {pct(to_float(finv_only) / to_float(diff_total) if diff_total else None)}，差异主要集中在转账类、高频消费类和部分负债类，其中 Rent 作为支出类别需要进行专项关注。",
        styles["body"],
    ))
    add_callout(story, "本段为报告的核心结论，读者可仅通过本段掌握样本范围、两侧覆盖率、双方非空一致率和主要差异来源。", styles, background=YELLOW_PALE, edge=ORANGE)
    story.append(PageBreak())

    # Part 2: Difference structure panorama
    add_section_header(story, "2. 差异结构全景", styles)
    story.append(paragraph("本部分先展示交易层面的差异类型分布，再通过业务板块交叉矩阵定位差异集中区域，最后列示 Top 差异流向。", styles["body"]))
    story.append(paragraph("2.1 总体差异类型分布", styles["h2"]))
    diff_table = table([
        ["差异类型", "数量", "占差异总数", "分析含义"],
        ["双方分类不一致", num(mismatch), pct(to_float(mismatch) / to_float(diff_total) if diff_total else None), "两侧均有分类，但类别标签不同"],
        ["仅 Illion 有值", num(illion_only), pct(to_float(illion_only) / to_float(diff_total) if diff_total else None), "finv 侧缺少分类"],
        ["仅 finv 有值", num(finv_only), pct(to_float(finv_only) / to_float(diff_total) if diff_total else None), "finv 新增识别或 Illion 漏识别"],
        ["总差异数", num(diff_total), "100.00%", "分类不一致与单边缺失的合计"],
        ["双方均为空", num(both_empty), "不纳入", "不进入差异流向和类别差异分析"],
    ], [38 * mm, 22 * mm, 25 * mm, 85 * mm], font_size=7.8)
    story.append(diff_table)
    add_callout(story, f"仅 finv 有值共 {num(finv_only)} 笔，占差异总数 {pct(to_float(finv_only) / to_float(diff_total) if diff_total else None)}，是当前最大的差异来源。", styles, background=YELLOW_PALE, edge=ORANGE)

    story.append(paragraph("2.2 业务板块定义", styles["h2"]))
    grouping_rows = [["业务板块", "类别数", "类别范围", "处理方式"]]
    for group in REPORT_GROUP_ORDER:
        grouping_rows.append([
            REPORT_GROUP_LABELS[group], len(REPORT_GROUPS[group]), ", ".join(REPORT_GROUPS[group]),
            "转账不计入收入或支出" if group == "transfer" else ("保留指标摘要，不重复 Liability 专项分析" if group == "liability" else ("Rent 重点分析" if group == "expense" else "完整分析")),
        ])
    story.append(table(grouping_rows, [24 * mm, 14 * mm, 94 * mm, 38 * mm], font_size=6.8))
    story.append(Spacer(1, 5))

    story.append(paragraph("2.3 业务板块差异汇总", styles["h2"]))
    summary_rows = [["业务板块", "Illion侧涉及差异", "finv侧涉及差异", "板块内冲突", "Illion源跨板块", "finv目标跨板块", "仅 Illion", "仅 finv"]]
    for group in REPORT_GROUP_ORDER:
        value = block_analysis["summary"][group]
        summary_rows.append([
            REPORT_GROUP_LABELS[group], num(value["illion_side"]), num(value["finv_side"]), num(value["within"]),
            num(value["source_cross"]), num(value["target_cross"]), num(value["illion_only"]), num(value["finv_only"]),
        ])
    story.append(table(summary_rows, [24 * mm, 22 * mm, 22 * mm, 20 * mm, 24 * mm, 24 * mm, 18 * mm, 18 * mm], font_size=6.7))
    story.append(paragraph("注：板块侧涉及差异为源类别或目标类别视角，可能对同一笔跨板块差异分别计入两侧；板块交叉矩阵为互斥口径，矩阵单元合计对应差异明细。", styles["small"]))
    story.append(PageBreak())

    story.append(paragraph("2.4 业务板块交叉矩阵", styles["h2"]))
    matrix_rows = [["Illion板块 / finv板块"] + [REPORT_GROUP_LABELS.get(col, col) for col in REPORT_GROUP_ORDER] + ["finv为空"]]
    for row_group in REPORT_GROUP_ORDER + ["Illion为空"]:
        row_label = REPORT_GROUP_LABELS.get(row_group, row_group)
        matrix_rows.append([row_label] + [num(block_analysis["matrix"][(row_group, col)]) for col in REPORT_GROUP_ORDER + ["finv为空"]])
    story.append(table(matrix_rows, [34 * mm, 27 * mm, 27 * mm, 27 * mm, 27 * mm, 27 * mm], font_size=7.4))
    story.append(paragraph("矩阵中的对角线表示同一业务板块内的类别冲突；非对角线表示跨板块迁移；最后一列和最后一行分别表示单边缺失。", styles["body"]))

    story.append(paragraph("2.5 Top 差异流向", styles["h2"]))
    flow_rows = [["排名", "Illion类别", "finv类别", "差异类型", "数量", "差异金额", "主要方向"]]
    for index, flow in enumerate(flows[:20], 1):
        flow_rows.append([
            index, safe_text(flow.get("illion_category")), safe_text(flow.get("finv_category")), safe_text(flow.get("diff_type")),
            num(flow.get("count")), money(flow.get("diff_amount")), safe_text(flow.get("direction")),
        ])
    story.append(table(flow_rows, [10 * mm, 29 * mm, 29 * mm, 26 * mm, 16 * mm, 24 * mm, 30 * mm], font_size=6.7))
    story.append(PageBreak())

    # Part 3: Business block category deep dives
    add_section_header(story, "3. 分业务板块的类型细探", styles)
    story.append(paragraph(f"本部分对 {category_count} 个类别统一展示七项核心指标：Illion 覆盖率、finv 覆盖率、并集数量、交集数量、交集占比（并集）、Illion 独有占比和 finv 独有占比。重点类别补充差异流向；底稿没有提供的字段留空。", styles["body"]))

    story.append(paragraph("3.1 收入类", styles["h2"]))
    story.append(paragraph("收入类包括 Wages、Centrelink 和 All Other Credits，不包括转账类。重点观察收入与转账之间的分类边界。", styles["body"]))
    for item in categories:
        if report_group(item.get("category")) == "income":
            add_category_block_v2(story, item, flows, styles, detailed=True)

    story.append(PageBreak())
    story.append(paragraph("3.2 支出类", styles["h2"]))
    story.append(paragraph("支出类共 23 个类别，所有类别使用同一套七项指标；Rent 单独重点分析，高频消费类别补充 Top 差异流向。", styles["body"]))
    for item in categories:
        if report_group(item.get("category")) == "expense":
            add_category_block_v2(story, item, flows, styles, detailed=item.get("category") in REPORT_FOCUS)

    story.append(PageBreak())
    rent = next((item for item in categories if item.get("category") == "Rent"), None)
    if rent:
        story.append(paragraph("3.2.1 Rent 专项", styles["h2"]))
        story.append(paragraph("Rent 重点观察其七项核心指标，以及与 External Transfers、Internal Transfer、Utilities 和 Home Improvement 的差异流向。", styles["body"]))
        add_category_block_v2(story, rent, flows, styles, detailed=True)

    story.append(PageBreak())
    story.append(paragraph("3.3 转账类", styles["h2"]))
    story.append(paragraph("External Transfers 与 Internal Transfer 作为独立板块分析，不纳入收入或支出合计；转账类额外关注 credit / debit 方向。底稿未提供的完整方向汇总留空。", styles["body"]))
    for item in categories:
        if report_group(item.get("category")) == "transfer":
            add_category_block_v2(story, item, flows, styles, detailed=True)
    story.append(Spacer(1, 5))
    add_callout(story, "转账类处理原则：纳入总体覆盖率、差异率和分类迁移分析，但不计入收入合计、支出合计或净收入计算。", styles)

    story.append(PageBreak())
    story.append(paragraph("3.4 负债类摘要", styles["h2"]))
    story.append(paragraph("负债类保留 8 个类别的七项核心指标，作为全分类报告的一部分；贷款生命周期、Counterparty matching、Dishonours 和 Unknown Loans 的深度分析引用已有 Liability 专项报告，本报告不重复展开。", styles["body"]))
    for item in categories:
        if report_group(item.get("category")) == "liability":
            add_category_block_v2(story, item, flows, styles, detailed=False)
    add_callout(story, "Liability 专项分析之外，本报告只关注负债类别在整体差异结构中的位置和七项类别指标。", styles, background=YELLOW_PALE, edge=ORANGE)

    story.append(PageBreak())
    story.append(paragraph("3.5 跨板块观察与建议", styles["h2"]))
    story.append(paragraph("跨板块优化建议应优先围绕转账与收入、转账与 Rent、高频消费类别边界以及 finv 独有分类的合理性验证展开。具体优先级以 Top 差异流向和类别七项指标共同判断。", styles["body"]))
    action_rows = [
        ["观察方向", "建议动作", "数据情况"],
        ["覆盖范围差异", "核验 finv 独有分类，区分合理扩展与误归类", "底稿可提供类别数量和差异流向；完整人工判定留空"],
        ["转账边界", "重点检查转账与 Wages、All Other Credits、Rent 的迁移", "底稿提供差异流向和主要方向"],
        ["高频消费边界", "重点检查 Groceries、Dining Out、Retail、Automotive、Gambling", "底稿提供类别指标和 Top 流向"],
        ["Rent", "单独建立 Rent 规则回归集", "底稿提供类别指标和部分差异流向"],
        ["负债类", "引用 Liability 专项报告，不重复展开", "本报告仅保留类别指标摘要"],
    ]
    story.append(table(action_rows, [32 * mm, 78 * mm, 60 * mm], font_size=7.2))

    # Appendices
    story.append(PageBreak())
    add_section_header(story, f"附录 A：{category_count} 类别七项指标完整表", styles)
    story.append(paragraph("所有类别均按相同口径展示。底稿中不存在的字段保持为空。", styles["body"]))
    score_rows = [["类别", "业务板块", "Illion覆盖率", "finv覆盖率", "并集", "交集", "交集占比", "Illion独有占比", "finv独有占比"]]
    for item in categories:
        score_rows.append([
            item.get("category"), REPORT_GROUP_LABELS.get(report_group(item.get("category")), ""),
            pct(item.get("illion_coverage")), pct(item.get("finv_coverage")), num(item.get("union_count")),
            num(item.get("intersection_count")), pct(item.get("intersection_share")), pct(item.get("illion_only_share")),
            pct(item.get("finv_only_share")),
        ])
    story.append(table(score_rows, [26 * mm, 20 * mm, 18 * mm, 18 * mm, 14 * mm, 14 * mm, 19 * mm, 20 * mm, 20 * mm], font_size=6.1))

    story.append(PageBreak())
    add_section_header(story, "附录 B：指标定义与数据缺失说明", styles)
    definition_rows = [
        ["指标", "定义"],
        ["Illion 覆盖率", "Illion 该类别数量 / 总交易数"],
        ["finv 覆盖率", "finv 该类别数量 / 总交易数"],
        ["并集数量", "至少一方识别为该类别的交易数"],
        ["交集数量", "双方均识别为该类别的交易数"],
        ["交集占比（并集）", "交集数量 / 并集数量"],
        ["Illion 独有占比", "Illion 独有数量 / 并集数量"],
        ["finv 独有占比", "finv 独有数量 / 并集数量"],
    ]
    story.append(table(definition_rows, [45 * mm, 125 * mm], font_size=7.6))
    story.append(Spacer(1, 8))
    story.append(paragraph("数据缺失处理：本报告仅使用数据底稿已有字段和由这些字段直接计算得到的汇总指标。底稿未提供的板块金额、完整方向汇总、人工准确性判定和完整案例结论不进行推断，在报告中留空或不展示。", styles["body"]))

    doc.build(story)
    return {"output": str(output_path), "categories": len(categories), "details": len(details), "flows": len(flows), "metrics": metrics}


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate Chinese full category performance report PDF")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    result = build_report_v2(args.input, args.output)
    print(f"Generated: {result['output']}")
    print(f"Categories: {result['categories']}; detail rows: {result['details']}; top flows: {result['flows']}")


if __name__ == "__main__":
    main()
